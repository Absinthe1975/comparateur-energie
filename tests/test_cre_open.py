import datetime, io, json, os, shutil, sys, tempfile, unittest
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))
import cre_open as c, update_prices as up
import openpyxl
DEPOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.join(DEPOT, "donnees")

def xlsx():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Historique grilles"
    for d, k0, k1 in ((datetime.datetime(2026, 10, 1), 0.17832, 0.14557), (datetime.datetime(2026, 9, 1), 0.16757, 0.13474)):
        ws.append([d, "Cuisson/eau chaude", None, "Chauffage"]); ws.append([None, "HT", "TTC", "HT", "TTC"])
        ws.append(["Abonnement (€/an)", 112.8, 152.4596, 252.04, 360.7916]); ws.append(["Part variable moyenne (€/kWh)", k0 - .05, k0, k1 - .04, k1]); ws.append([])
    h = wb.create_sheet("Historique PRVG moyen"); h.append(["Date", "HT", "TTC"])
    for d, v in ((datetime.datetime(2026, 9, 1), 172.05), (datetime.datetime(2026, 10, 1), 182.875), (datetime.datetime(2026, 11, 1), 194.9)): h.append([d, v / 1.4, v])
    g = wb.create_sheet("Grille tarifaire TTC"); g.append(["Grilles tarifaires TTC prix repère au 1er novembre 2026"])
    g.append([None, "Consommateur type", "Part fixe", "Minimale", "Moyenne", "Maximale"])
    g.append(["Cuisson/eau chaude", 152.46, 185.6, 190.27, 200.0]); g.append(["Chauffage", 360.79, 147.98, 157.6, 176.39])
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def rows(fin=None, kwh=0.2001):
    return {"data": [{"DATE_DEBUT": "08/01/2026", "DATE_FIN": fin, "P_SOUSCRITE": 6, "PART_FIXE_HT": 143.28, "PART_FIXE_TTC": 190.81, "PART_VARIABLE_HT": 0.1361, "PART_VARIABLE_TTC": kwh},
                     {"DATE_DEBUT": "02/01/2026", "DATE_FIN": "31/07/2026", "P_SOUSCRITE": 6, "PART_FIXE_HT": 141.6, "PART_FIXE_TTC": 188.2, "PART_VARIABLE_HT": 0.1308, "PART_VARIABLE_TTC": 0.194}]}

class T(unittest.TestCase):
    def test_gaz_octobre_puis_novembre(self):
        g = c.gaz_grilles(xlsx())
        d, x = c.gaz_en_vigueur(g, datetime.date(2026, 10, 13)); self.assertEqual((d.month, x["b1_kwh"], x["b0_kwh"], x["b1_abo"]), (10, 0.1456, 0.1783, 360.79))
        d, x = c.gaz_en_vigueur(g, datetime.date(2026, 11, 2)); self.assertEqual((d.month, x["b1_kwh"], x["b0_kwh"]), (11, 0.1576, 0.1903))
    def test_trv_sans_se_fier_aux_dates(self):
        self.assertEqual(c.trv_courant(rows()["data"], 6)["PART_VARIABLE_TTC"], 0.2001)
        self.assertEqual(c.trv_courant(rows("nul", "0,2001")["data"], 6)["PART_VARIABLE_TTC"], 0.2001)
    def test_trv_refuse_zero_ou_deux_lignes_en_vigueur(self):
        with self.assertRaises(ValueError): c.trv_courant(rows("31/07/2026")["data"], 6)
        r = rows()["data"]; r[1]["DATE_FIN"] = None
        with self.assertRaises(ValueError): c.trv_courant(r, 6)
    def test_trv_refuse_prix_incoherent(self):
        with self.assertRaises(ValueError): c.trv_courant(rows(None, 0.9)["data"], 6)
    def test_repli_csv_quand_api_en_panne(self):
        def get(url):
            if "tabular-api" in url: raise RuntimeError("503")
            return "DATE_DEBUT;DATE_FIN;P_SOUSCRITE;PART_FIXE_HT;PART_FIXE_TTC;PART_VARIABLE_HT;PART_VARIABLE_TTC\n01/08/2026;;6;143,28;190,81;0,1361;0,2001\n".encode()
        t, how = c.trv("rid", 6, get); self.assertEqual((t["PART_VARIABLE_TTC"], how), (0.2001, "trv_depuis_csv"))

    def test_variation_officielle(self):
        v = c.variation_mensuelle(c.prvg_moyen(xlsx()), datetime.date(2026, 10, 13)); self.assertAlmostEqual(v, 6.29, 1)
        v = c.variation_mensuelle(c.prvg_moyen(xlsx()), datetime.date(2026, 11, 2)); self.assertAlmostEqual(v, 6.57, 1)

class Integ(unittest.TestCase):
    def run_official(self, today, trv_kwh=0.2001, trv_abo_manquants=False):
        d = __import__("donnees").charger(ROOT)
        if trv_abo_manquants:
            for k in ("12", "15", "18", "24", "30", "36"): d["trv_abo"].pop(k, None)
        src = json.load(open(os.path.join(ROOT, "sources.json"), encoding="utf-8"))
        def get(url):
            if "tabular-api" in url:
                kva = int(url.split("exact=")[1].split("&")[0])
                if kva == 6: return json.dumps(rows(None, trv_kwh)).encode()
                return json.dumps({"data": [dict(rows()["data"][0], P_SOUSCRITE=kva, PART_VARIABLE_TTC=0.1985, PART_FIXE_TTC=239.84 if kva == 9 else round(25.99 * 12 * kva / 12, 2))]}).encode()
            if url.endswith(".html"): return b'<a href="/content/download/1/fichier.xlsx">x</a>'
            return xlsx()
        class N: pass
        now = N(); now.date = lambda: today
        rep = up.Report(); up.update_official(d, src, rep, now, get); return d, rep
    def test_octobre_confirme_nos_valeurs(self):
        d, rep = self.run_official(datetime.date(2026, 10, 13)); self.assertEqual(rep.failed, []); self.assertEqual(d["gas_b1"]["a"], 0.1456); self.assertEqual(d["gas_b0"]["a"], 0.1783)
    def test_novembre_applique_la_nouvelle_grille(self):
        d, rep = self.run_official(datetime.date(2026, 11, 2)); self.assertEqual((d["gas_b1"]["a"], d["gas_ref"]["a"], d["gas_b0"]["a"]), (0.1576, 0.1576, 0.1903)); self.assertTrue(rep.applied)
    def test_abonnements_des_autres_puissances(self):
        d, rep = self.run_official(datetime.date(2026, 10, 13), trv_abo_manquants=True); self.assertEqual(rep.failed, [])
        for k in ("12", "15", "18", "24", "30", "36"): self.assertIn(k, d["trv_abo"])
        self.assertEqual(d["trv_abo"]["12"], round(25.99 * 12, 2))
    def test_abonnements_saisis_non_ecrases_par_la_cre(self):
        d, rep = self.run_official(datetime.date(2026, 10, 13)); self.assertEqual(d["trv_abo"]["12"], 285.12)
        self.assertTrue(any("tarif réglementé 12 kVA" in x and "à vérifier" in x for x in rep.review))
    def test_prix_aberrant_non_applique(self):
        d, rep = self.run_official(datetime.date(2026, 10, 13), 0.3); self.assertEqual(d["elec_ref"]["a"], 0.2001); self.assertTrue(any("hors limites" in x for x in rep.review))

if __name__ == "__main__": unittest.main()
